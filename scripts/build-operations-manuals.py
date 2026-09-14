from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "operations" / "manuals"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
PALE = "EFF4F8"
GRAY = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), GRAY)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.color.rgb = BLACK
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def configure(doc, code, title):
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (
        ("Title", 27, 0, 14),
        ("Heading 1", 18, 18, 8),
        ("Heading 2", 13.5, 13, 5),
        ("Heading 3", 11, 9, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    # Remove theme borders so title pages remain clean and printer friendly.
    for name in ("Title", "Subtitle"):
        p_pr = styles[name]._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)
    styles["Caption"].font.color.rgb = BLACK
    for header in (section.header, section.even_page_header, section.first_page_header):
        hp = header.paragraphs[0]
        hp.text = ""
        hp.style = styles["Caption"]
    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        fp = footer.paragraphs[0]
        fp.text = f"{code}   Controlled draft   "
        fp.style = styles["Caption"]
        page_field(fp)
    doc.core_properties.title = title
    doc.core_properties.subject = "Business operations and standard operating procedures"
    doc.core_properties.author = "NexPhase Labs"


def add_title(doc, title, subtitle, code, revision="Revision 0 1"):
    p = doc.add_paragraph(title, style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub = doc.add_paragraph(subtitle)
    sub.style = doc.styles["Subtitle"]
    for run in sub.runs:
        run.font.color.rgb = BLACK
    add_table(
        doc,
        ["Document", "Revision", "Status", "Review cycle"],
        [[code, revision, "Controlled draft", "Annual and after material change"]],
        [1.25, 1.15, 1.7, 2.35],
    )
    p = doc.add_paragraph()
    r = p.add_run("Issuance condition  ")
    r.bold = True
    p.add_run(
        "The authorized members must approve the authority matrix and open policies, then named operators must complete a staging walkthrough. Record approval and training evidence in Operating Controls before changing this document to issued."
    )


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 16 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(str(text))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        cant_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            if row_idx % 2:
                shade(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[idx])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(headers[idx]) < 12 and len(str(value)) < 25 else WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(str(value))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.add_run(item)


def commands(doc, lines):
    """Literal commands, monospaced and shaded so they are copied exactly.

    16.6: MAN 002 described the deployment and recovery procedures rather than
    stating them. A description is not something a second person can follow at
    two in the morning. Anything in one of these blocks is typed verbatim.
    """
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.18)
        run = p.add_run(line)
        run.font.name = "Cascadia Mono"
        run.font.size = Pt(9)
        r_pr = run._r.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
            r_fonts.set(qn(attribute), "Cascadia Mono")
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), PALE)
        r_pr.append(shd)


def section(doc, title, paragraphs=None, bullets_list=None):
    doc.add_heading(title, level=1)
    for text in paragraphs or []:
        doc.add_paragraph(text)
    if bullets_list:
        bullets(doc, bullets_list)


def sop_header(doc, number, title, owner, purpose, scope, prerequisites, records, metric):
    doc.add_heading(f"{number} {title}", level=1)
    add_table(
        doc,
        ["Owner", "Purpose", "Scope"],
        [[owner, purpose, scope]],
        [1.3, 2.6, 2.85],
    )
    doc.add_heading("Prerequisites", level=2)
    bullets(doc, prerequisites)
    doc.add_heading("Responsibilities", level=2)
    doc.add_paragraph(
        "The owner performs or assigns the work. The current assignee records each action and exception. An administrator reviews controlled closure or readiness decisions. Quality personnel decide lot disposition. Commercial and legal decisions remain with their authorized owners."
    )
    doc.add_heading("Procedure", level=2)


def sop_steps(doc, rows):
    add_table(doc, ["Step", "Action", "Expected record"], rows, [0.55, 3.65, 2.55])


def sop_close(doc, exceptions, completion, metric, records):
    doc.add_heading("Exceptions and escalation", level=2)
    bullets(doc, exceptions)
    doc.add_heading("Completion criteria", level=2)
    bullets(doc, completion)
    doc.add_heading("Measures and records", level=2)
    add_table(doc, ["Measure", "Required records"], [[metric, records]], [3.25, 3.5])


def build_master():
    doc = Document()
    configure(doc, "MAN 001", "Business Operations Control Manual")
    add_title(doc, "Business Operations Control Manual", "How NexPhase assigns work controls evidence and approves readiness", "MAN 001")
    doc.add_heading("Manual purpose", level=1)
    doc.add_paragraph(
        "This manual defines the operating system around NexPhase transactions. It tells staff where work begins, who owns the next action, what evidence completes it, how exceptions are escalated, and when the business may accept real orders. The application records operational facts; controlled files hold source evidence and issued documents."
    )
    doc.add_heading("Operating status", level=1)
    doc.add_paragraph(
        "The transaction engine is built and automated tests cover the current workflows. Production acceptance remains closed until every launch critical control has a named owner, current evidence, and administrator review. Simulated provider results do not prove payment, tax, email, or carrier readiness."
    )
    add_table(doc, ["Term", "Meaning"], [
        ["Built", "The workflow and its data rules exist in the application."],
        ["Verified", "Automated tests or a controlled rehearsal demonstrate the expected path and failure path."],
        ["Operationally ready", "A named owner, approved rule, real data or configuration, evidence, exception path, and monitoring exist."],
        ["Launch approved", "Authorized members accept the remaining business risk and record the decision."],
    ], [1.55, 5.2])
    section(doc, "Operating principles", bullets_list=[
        "No material becomes sellable without a named quality disposition.",
        "No control becomes ready from a passing software check alone.",
        "One person owns each current action; handoffs include a deadline and note.",
        "Corrections create attributed history. Staff do not edit records outside controlled workflows.",
        "Real provider activity begins only after sandbox evidence, approved policy, and production authorization.",
        "A carrier label remains active until its refund is confirmed. Staff never retry an uncertain label purchase or refund request.",
        "An operational exception is opened as a case when routine workflow cannot safely resolve it.",
    ])
    doc.add_heading("Authority and responsibilities", level=1)
    add_table(doc, ["Person", "Operating responsibility", "Reserved decisions"], [
        ["Mel", "Brand, supply, product, quality, lot decisions", "Exceptional release requires Sam countersignature when the approved trigger applies"],
        ["Sam", "Sales, terms, pricing, customer approval, fulfillment, banking, signatures", "Commercial acceptance, banking, signature, and launch decisions assigned by the members"],
        ["Tima", "Daily order queue, customer communication, verification preparation, supplier document follow up, contract first review", "Prepares records and escalates decisions outside delegated authority"],
        ["Ammre", "Platform, infrastructure, controls, bookkeeping, payment and vendor analysis", "No signing, catalog publication, or lot release authority unless separately approved"],
        ["Administrator", "Account administration and controlled system decisions", "Records readiness and case closure after reviewing evidence"],
    ], [1.0, 3.0, 2.75])
    doc.add_paragraph(
        "The application currently has admin, quality, and operations roles. Before issue, the members must map every named person to the exact application permissions and approve the exceptional release trigger. Until then, role access is a technical safeguard and does not expand business authority."
    )
    doc.add_heading("Systems of record", level=1)
    add_table(doc, ["Record", "Authoritative location", "Rule"], [
        ["Suppliers, lots, inventory, orders, payments, shipments, returns", "NexPhase application", "Use controlled forms and append only history. Do not maintain a duplicate tracker."],
        ["Operating controls and cases", "NexPhase application", "Assign one owner, date the next action, link evidence, and preserve history."],
        ["Signed agreements, source evidence, issued SOPs, training", "Approved Shared Drive", "Restrict by role. Link from the application where a record requires evidence."],
        ["Setup and project actions", "Operations tracker", "Store tasks and canonical links only. Do not copy live quantities or status."],
        ["Credentials", "Approved secret and deployment systems", "Never place secrets in Drive notes, manuals, or application comments."],
    ], [1.8, 1.7, 3.25])
    doc.add_heading("Daily operating rhythm", level=1)
    sop_steps(doc, [
        ["1", "Open Today’s Operations and review urgent counts.", "Prioritized work list"],
        ["2", "Open Orders assigned to me, overdue service actions, and unassigned active orders.", "Every actionable order has one owner and future due time"],
        ["3", "Review operational cases assigned to me, critical cases, and overdue cases.", "Containment and next actions are current"],
        ["4", "Review lots in quarantine and on hold; perform only role-authorized actions.", "Lot status and inventory ledger agree"],
        ["5", "Review payment, refund, notification, and procurement queues assigned to the role.", "Exceptions are assigned or opened as cases"],
        ["6", "Update assigned operating controls and link evidence when work is ready for review.", "No unsupported ready status"],
        ["7", "Before sign out, hand off unfinished orders and cases with a date and clear note.", "No hidden work or ownerless handoff"],
    ])
    doc.add_heading("Weekly review", level=1)
    bullets(doc, [
        "Review overdue orders, cases, operating controls, purchase orders, supplier reviews, retest dates, refunds, and failed notifications.",
        "Review inventory movements against physical count evidence and investigate variances through an operational case.",
        "Review active staff, role fit, one time passwords, and leavers. Record the access review as control evidence.",
        "Review provider and shipping exceptions. Do not treat simulated results as live evidence.",
        "Record decisions, owners, dates, and evidence links before ending the review.",
    ])
    doc.add_heading("Operating controls", level=1)
    add_table(doc, ["Status", "Use"], [
        ["Not started", "No work or evidence has been recorded."],
        ["In progress", "An active owner is working toward a dated result."],
        ["Blocked", "A note identifies the blocker and the owner keeps the due date current."],
        ["Awaiting review", "The owner linked completion evidence for administrator review."],
        ["Ready", "An administrator reviewed the owner, evidence, and applicable acceptance criteria."],
        ["Not applicable", "An administrator recorded the waiver reason. This is not a shortcut around missing evidence."],
    ], [1.5, 5.25])
    doc.add_heading("Operational cases", level=1)
    doc.add_paragraph(
        "Open a case for a complaint, deviation, supplier issue, incident, corrective or preventive action, or recall. High and critical cases require containment before they progress. A recall requires the linked lot to be on hold or withdrawn. An administrator closes a case only after root cause, corrective action, evidence, effectiveness review, and closure summary are complete."
    )
    doc.add_heading("Escalation levels", level=1)
    add_table(doc, ["Severity", "Response"], [
        ["Low", "Owner records and resolves within normal queue management."],
        ["Medium", "Owner sets a prompt due date and informs the responsible function."],
        ["High", "Contain immediately, notify the responsible member, and review related lots or orders."],
        ["Critical", "Stop affected activity, preserve evidence, notify authorized leadership immediately, and open an incident or recall case."],
    ], [1.25, 5.5])
    doc.add_heading("Operational measures", level=1)
    add_table(doc, ["Measure", "Review"], [
        ["Launch critical controls open and overdue", "Daily before production acceptance; weekly thereafter"],
        ["Orders unassigned, overdue, awaiting payment, fulfilling, refund due", "Daily"],
        ["Cases open, critical, overdue, and assigned to each owner", "Daily and in weekly review"],
        ["Lots in quarantine, on hold, withdrawn, exhausted, or nearing retest", "Daily for active work; weekly trend"],
        ["Inventory movement variances and physical count differences", "At every count and monthly close"],
        ["Provider failures, failed notices, label voids, refund exceptions", "Daily and at close"],
    ], [3.5, 3.25])
    doc.add_heading("Document control and training", level=1)
    bullets(doc, [
        "The document owner proposes revisions after application or policy changes.",
        "The approver compares the procedure to the verified interface and records approval, effective date, and next review.",
        "Each operator completes a walkthrough using staging and records questions or corrections before issue.",
        "The controlled copy is stored in the approved Shared Drive. Superseded copies remain identifiable and unavailable for routine use.",
        "A material interface, authority, legal, safety, provider, or workflow change triggers immediate review.",
    ])
    doc.add_heading("Approval record", level=1)
    add_table(doc, ["Role", "Name", "Decision", "Date", "Evidence"], [["Document owner", "", "", "", ""], ["Operations approver", "", "", "", ""], ["Quality approver", "", "", "", ""], ["Member approval", "", "", "", ""]], [1.45, 1.25, 1.15, 1.0, 1.9])
    path = OUT / "MAN-001_Business_Operations_Control_Manual.docx"
    doc.save(path)
    return path


SOPS = [
    {
        "number": "SOP 003", "title": "Daily Queue and Work Handoff", "owner": "Operations lead",
        "purpose": "Keep active work visible, owned, dated, and attributable.", "scope": "Orders, cases, controls, lots, procurement, notices, and refunds.",
        "prereq": ["Active staff account with an approved role", "Today’s Operations dashboard available", "Current escalation contacts"],
        "steps": [
            ["1", "Open Today’s Operations and note urgent or overdue counts.", "Start of day review recorded when required"],
            ["2", "Open Orders and select Assigned to me, Overdue service actions, then Unassigned active orders.", "Every active order has an appropriate owner"],
            ["3", "Claim only work you will own. Set a future next action due time and describe the next action.", "Order event records owner, due time, note, and actor"],
            ["4", "Open Cases and review assigned, critical, and overdue records. Update containment and next action.", "Case status and evidence match current facts"],
            ["5", "Review quarantine, hold, payment, refund, notification, and purchasing queues relevant to your role.", "Exceptions are acted on or assigned"],
            ["6", "Hand unfinished work to an active teammate with a due time and note before sign out.", "No hidden or ambiguous handoff"],
        ],
        "exceptions": ["Do not reassign work owned by another person unless you are the current owner or administrator.", "Open an incident case when the application is unavailable or the queue may be incomplete.", "Escalate any critical, safety, legal, payment, or release issue immediately."],
        "completion": ["Every actionable order and case has one current owner and current due date.", "Overdue work has an update, escalation, or documented blocker.", "All handoffs appear in append only history."],
        "metric": "Unassigned active orders, overdue orders, overdue cases, and handoffs without notes", "records": "Order events, case events, dashboard counts, and operating control evidence",
    },
    {
        "number": "SOP 004", "title": "Supplier Qualification and Purchasing", "owner": "Supply owner and purchasing operator",
        "purpose": "Buy only from an active supplier with a current, reviewable qualification.", "scope": "Supplier onboarding, qualification, purchase orders, expected receipts, and suspension.",
        "prereq": ["Approved qualification criteria and decision authority", "Controlled supplier evidence folder", "Approved products, costs, terms, and purchase authority"],
        "steps": [
            ["1", "Create or update the supplier identity, address, contacts, website, and internal notes.", "Attributed supplier record"],
            ["2", "Review identity, manufacturer or facility information, quality records, references, commercial terms, and applicable risks outside the application.", "Controlled evidence set"],
            ["3", "Record qualification reason, approved supply scope, evidence link, and future review date.", "Named qualification event"],
            ["4", "Create a draft purchase order with product, quantity, line cost, freight, duty, dates, and supplier reference.", "Draft PO and line records"],
            ["5", "Review approval threshold and issued document requirements. When authorized, send the order and record the transition.", "Sent PO event and supplier acknowledgment evidence"],
            ["6", "Receive material only against the matching open line when applicable. Investigate differences before closing short.", "Expected receipt links to actual lots"],
            ["7", "Suspend a supplier immediately when continued purchasing is not approved. Record the reason.", "Suspension event blocks new POs"],
        ],
        "exceptions": ["Do not qualify without scope, evidence link, and review date.", "Do not raise a PO for an unqualified, suspended, inactive, or changed supplier.", "Open a supplier issue case for material quality, identity, delivery, or documentation failures."],
        "completion": ["Qualification is current and within scope.", "PO totals and expected receipts are complete and attributable.", "Supplier acknowledgment and invoice matching evidence are linked before financial close."],
        "metric": "Qualifications overdue, unacknowledged sent POs, overdue receipts, and supplier issue cases", "records": "Supplier record and events, evidence folder, PO and events, lots, supplier invoices, reconciliation",
    },
    {
        "number": "SOP 005", "title": "Lot Receipt and Quarantine", "owner": "Receiving operator",
        "purpose": "Record received material without making it sellable.", "scope": "Physical receipt, count, condition, purchase order match, quarantine, and initial documents.",
        "prereq": ["Active supplier and expected receipt when a PO exists", "Approved receiving area and labels", "Calibrated tools and current safety information"],
        "steps": [
            ["1", "Inspect package identity, damage, tamper evidence, temperature condition, and accompanying documents before opening.", "Receiving inspection evidence"],
            ["2", "Count or measure independently and compare supplier document, package, and expected receipt.", "Verified received quantity"],
            ["3", "Create the lot with exact lot number, product, received quantity, date, supplier, manufacturer information, location, and landed cost when authorized.", "Lot in quarantine and receipt movement"],
            ["4", "Apply the quarantine label and segregate the material from released inventory.", "Physical status matches application status"],
            ["5", "Attach the available COA or controlled source documents. Do not treat upload as release.", "Document history linked to lot"],
            ["6", "Open a deviation or supplier issue case for damage, mismatch, missing identity, or count variance.", "Exception owner and containment"],
        ],
        "exceptions": ["Do not receive against a closed or stale expected receipt.", "Do not change received quantity outside the controlled correction or adjustment workflow.", "Hold damaged or ambiguous material and preserve packaging evidence."],
        "completion": ["Lot exists in quarantine with an attributed receipt movement.", "Physical label and segregated location match the record.", "Differences are resolved or assigned in a case."],
        "metric": "Receipt discrepancies, time to quarantine record, and lots missing documents or manufacturer identity", "records": "Lot, receipt movement, PO line, inspection evidence, lot documents, case when needed",
    },
    {
        "number": "SOP 006", "title": "Lot Testing and Disposition", "owner": "Quality owner",
        "purpose": "Release only material that meets approved specifications and evidence requirements.", "scope": "Test results, COA review, labels, release, hold, withdrawal, and exhaustion.",
        "prereq": ["Approved specifications, methods, and reviewer competence", "Lot is in quarantine or controlled hold", "Manufacturer identity and required documents are available"],
        "steps": [
            ["1", "Verify lot identity, manufacturer name and address, received quantity, source documents, and current specification version.", "Complete review set"],
            ["2", "Record analytical results with method, units, limits, result, date, and named recorder.", "Attributed test result history"],
            ["3", "Review every release blocker shown by the application. Resolve blockers through corrected evidence or a case.", "No unresolved release blocker"],
            ["4", "Record the quality disposition. Release only when requirements pass; otherwise hold or withdraw with a reason.", "Named status event"],
            ["5", "Issue or retrieve the correct lot label and available COA only after the applicable status permits it.", "Controlled lot documents"],
            ["6", "Monitor retest and expiry requirements. Hold material before uncertainty affects a shipment.", "Current disposition"],
        ],
        "exceptions": ["Do not use an exceptional release until the members approve its exact trigger and countersignature rule.", "Put a released lot on hold before opening a recall.", "Open a deviation or CAPA when a test, document, or method failure needs investigation."],
        "completion": ["Disposition is attributable and physical status matches the application.", "Released lots have no release blockers.", "Hold and withdrawal reasons are clear and linked to cases when investigation is required."],
        "metric": "Quarantine age, release blockers, holds, withdrawals, retest alerts, and disposition reversals", "records": "Lot, tests, documents, status events, labels, COA, case evidence",
    },
    {
        "number": "SOP 007", "title": "Inventory Movements and Physical Counts", "owner": "Inventory operator",
        "purpose": "Keep physical stock and the lot ledger reconciled without side records.", "scope": "Receipt, sample, shipment, return, adjustment, destruction, count, and exhaustion.",
        "prereq": ["Lot identity and physical location confirmed", "Approved unit of measure", "Two distinct named witnesses for destruction"],
        "steps": [
            ["1", "Confirm the lot, current status, on hand quantity, unit, and active reservations.", "Correct movement source"],
            ["2", "Choose sample, destruction, or adjustment. Record quantity, occurrence date, and a specific reason.", "Validated movement input"],
            ["3", "For destruction, identify two different witnesses. For a positive adjustment, place released stock on hold first.", "Required control evidence"],
            ["4", "Submit the movement once. The application updates quantity and ledger in one transaction.", "Attributed ledger entry and new balance"],
            ["5", "For a physical count, compare every container to the application. Investigate the variance before recording an adjustment.", "Count sheet and case when unexplained"],
            ["6", "Review the movement report for the count period and confirm that no movement reduced stock below active reservations.", "Reconciled count evidence"],
        ],
        "exceptions": ["Do not use unlike units or a quantity larger than available unreserved stock.", "Do not increase a released lot or reopen an exhausted lot through an adjustment.", "Open a deviation case for an unexplained variance or failed destruction control."],
        "completion": ["Physical quantity equals the application balance or a case owns the difference.", "Every non sale movement names the actor, date, reason, and witnesses when required.", "Exhausted released lots have the exhausted status."],
        "metric": "Count accuracy, adjustment count and value, unexplained variances, and destruction exceptions", "records": "Count evidence, lot movements, lot status history, reservation review, case records",
    },
    {
        "number": "SOP 008", "title": "Operational Cases Recalls and Corrective Action", "owner": "Responsible function owner",
        "purpose": "Contain, investigate, correct, verify, and close operational exceptions.", "scope": "Complaints, deviations, supplier issues, incidents, CAPA, and recalls.",
        "prereq": ["Known facts and immediate safety or business containment", "Active owner and due date", "Related lot, order, supplier, communication, and evidence references"],
        "steps": [
            ["1", "Open the correct case type with severity, title, factual scope, owner, due date, and known links.", "Open case and creation event"],
            ["2", "For high or critical cases, stop affected activity and record containment before progressing.", "Contained status and action detail"],
            ["3", "For a recall, put the linked lot on hold or withdraw it first, then open the recall and identify affected orders or consignees.", "Controlled lot and recall scope"],
            ["4", "Investigate and record evidence, root cause, and the decision basis. Separate facts from assumptions.", "Investigation record"],
            ["5", "Record corrective action and preventive action, assign implementation, and preserve communication evidence.", "Action record and evidence link"],
            ["6", "Perform an effectiveness check after action. An administrator reviews the complete record and closure summary.", "Closed case and closure event"],
        ],
        "exceptions": ["Critical cases require immediate leadership notification outside the routine queue.", "Do not close while required action or effectiveness evidence is missing.", "Do not reopen a closed case except through administrator review."],
        "completion": ["Containment, root cause, correction, evidence, effectiveness, and closure are complete.", "Affected lots, orders, suppliers, and communications are linked.", "An administrator records closure after review."],
        "metric": "Cases by severity, containment time, overdue actions, recurrence, and effectiveness failures", "records": "Operational case and events, lot and order histories, evidence folder, communication log",
    },
    {
        "number": "SOP 009", "title": "Order Payment Fulfillment Shipping and Delivery", "owner": "Commercial and fulfillment owners",
        "purpose": "Move an accepted order through payment, lot allocation, documents, packing, shipment, and evidenced delivery with one current owner.", "scope": "Submitted through delivered orders, including USPS, UPS, and FedEx through the configured shipping provider.",
        "prereq": ["Approved customer eligibility and commercial terms", "Accepted server calculated quote and valid stock reservation", "Released lot, approved origin and parcel profile, and live provider evidence before production"],
        "steps": [
            ["1", "Claim or assign the order and set the next service deadline.", "Owner, due time, and handoff event"],
            ["2", "Confirm immutable order totals, accepted shipping quote, customer details, and payment status. Do not prepare on pending or failed payment.", "Order ready for fulfillment"],
            ["3", "Record or reconcile payment only from approved evidence. Resolve failed attempts before proceeding.", "Paid status and payment event"],
            ["4", "Select eligible released lots and begin preparation. Confirm the reservation is current.", "Lot allocation and fulfilling status"],
            ["5", "Perform packing checks, confirm the final dimensions and weight, and issue the packing slip and invoice when their blockers are clear.", "Packed parcel and issued documents"],
            ["6", "Choose the approved ship from location and retrieve configured USPS, UPS, and FedEx rates. Select only an allowed service and buy one label through the recorded quote.", "Origin specific rate quote and one active label"],
            ["7", "If the label is wrong and has not been tendered, enter the cancellation reason and request the refund once. Keep the parcel blocked while pending. Use Check refund status for recovery; never submit another provider refund request.", "Refund status and preserved label history"],
            ["8", "Create a replacement only after the prior label is confirmed voided. Record carrier handoff and tracking, move the order to shipped, and send the durable customer notice.", "Shipment movement, tracking, order event, notice"],
            ["9", "Monitor tracking and record delivery only from a carrier event, scan, signed receipt, or equivalent evidence. Queue the delivery notice. Open a case for loss, damage, or a shipment that misses the approved delivery threshold.", "Delivered date, evidence, order event, notice or open exception"],
        ],
        "exceptions": ["Do not use simulated rates, tax, payment, or email as production evidence.", "Do not ship a held, withdrawn, expired, exhausted, or inadequately reserved lot.", "Do not retry an uncertain label purchase or refund; use the app reconciliation control and compare the order history with Shippo.", "An active or unresolved label blocks order cancellation and replacement.", "Open a case for lost packages, material damage, dangerous goods uncertainty, duplicate provider records, or other provider inconsistencies."],
        "completion": ["Payment, lot, quantities, documents, selected origin, label history, tracking, delivery evidence, and order totals reconcile.", "Every cancelled label has a reason, requester, provider refund reference, and confirmed final status.", "Customer notification is delivered or assigned for recovery.", "The order has delivery evidence or an owned exception for the undelivered shipment."],
        "metric": "Time to owner, payment exceptions, pick and pack time, labels by origin, label failures, pending refunds, on time handoff, and delivery exceptions", "records": "Order and events, payment attempts, reservations, origin specific quotes, label and refund history, movements, issued documents, notices",
    },
    {
        "number": "SOP 010", "title": "Returns Refunds and Customer Complaints", "owner": "Customer service and finance owners",
        "purpose": "Resolve customer problems while keeping money, material, and communication records aligned.", "scope": "Feedback, returns, quarantined returned material, refund obligations, payment evidence, and complaint cases.",
        "prereq": ["Authenticated staff access", "Order and payment records", "Approved return, refund, and communication policy"],
        "steps": [
            ["1", "Acknowledge feedback and assign a current owner. Link the order when known.", "Owned feedback record"],
            ["2", "Classify quality, safety, delivery, documentation, commercial, or technical impact. Open a complaint case when investigation is required.", "Case and escalation path"],
            ["3", "Authorize a return only under approved policy. On receipt, inspect and quarantine returned material; never add it to sellable stock.", "Return event and quarantined movement"],
            ["4", "Calculate the recorded refund obligation from the order and returned lines. Obtain the required commercial or finance approval.", "Refund due amount"],
            ["5", "Send the refund through the approved rail, then record the actual amount and provider or bank reference.", "Refund evidence and reconciled status"],
            ["6", "Tell the customer the result through the controlled communication path and close feedback only when the case and money records agree.", "Communication and closure history"],
        ],
        "exceptions": ["Do not mark refunded before money has moved.", "Do not return received material to released inventory.", "Escalate safety, identity, contamination, repeat complaint, chargeback, or legal threat as a case."],
        "completion": ["Customer, order, return, refund, and case records agree.", "Refund reference and outstanding balance reconcile.", "Closure communication and evidence are attributable."],
        "metric": "First response time, overdue complaints, return cycle time, refund outstanding, recurrence, and chargebacks", "records": "Feedback, case, order events, return movement, refund reference, payment settlement, notices",
    },
    {
        "number": "SOP 011", "title": "Period Reporting and Month End Review", "owner": "Bookkeeping and finance reviewer",
        "purpose": "Produce period defined operational reports and reconcile them to external financial evidence.", "scope": "Orders, revenue, refunds, movements, inventory snapshot, settlements, AP, AR, tax, and close evidence.",
        "prereq": ["Approved close calendar and accounting mappings", "Bank and provider statements", "All period transactions, returns, refunds, receipts, and adjustments recorded"],
        "steps": [
            ["1", "Set the report From and Through dates. Confirm the period before downloading.", "Defined UTC reporting period"],
            ["2", "Download Orders and Movement Ledger CSVs. Save the current Inventory by Lot snapshot separately with its extraction time.", "Dated source exports"],
            ["3", "Reconcile order totals, shipping, refunds, outstanding refunds, and allocated costs to the application detail.", "Sales and refund reconciliation"],
            ["4", "Reconcile payment settlements and bank deposits to paid orders and recorded refunds.", "Provider and bank reconciliation"],
            ["5", "Match supplier invoices to POs and receipts. Record AP, AR, tax, inventory, and unresolved exceptions in the approved accounting system.", "Close schedules"],
            ["6", "Open a case for unexplained variance. Reviewer signs the close checklist only after exceptions are resolved or explicitly accepted.", "Reviewed close evidence"],
        ],
        "exceptions": ["Orders are selected by submission date; movements by occurrence date; inventory is a current snapshot.", "Do not treat application reports as a general ledger or tax filing system.", "Do not net unexplained differences into an adjustment without investigation and approval."],
        "completion": ["Exports and external statements reconcile to approved books.", "Refund, settlement, AP, AR, inventory, and tax exceptions have owners.", "Reviewer and completion date are recorded."],
        "metric": "Days to close, unreconciled items, refund outstanding, inventory variance, and late supplier invoices", "records": "Period CSVs, inventory snapshot, statements, reconciliation, case records, close checklist",
    },
    {
        "number": "SOP 012", "title": "Customer Notices and Service Recovery", "owner": "Customer service owner",
        "purpose": "Send attributable customer notices and recover failures without losing the required message.", "scope": "Order notices, feedback responses, delivery failures, bounced messages, and manual recovery.",
        "prereq": ["Approved sender and support inbox", "Approved message content and recipient", "Durable notification queue for order notices"],
        "steps": [
            ["1", "Review notifications needing attention and customer feedback assigned to you.", "Prioritized communication queue"],
            ["2", "Confirm recipient, order or case context, and permitted content before sending.", "Correct message target"],
            ["3", "Send through the approved channel. Order notices remain in the durable queue until delivered or handled.", "Delivery attempt history"],
            ["4", "Investigate provider errors, bounces, or missing configuration. Correct the cause and retry only when duplication risk is controlled.", "Recovery event"],
            ["5", "If another channel is used, record who contacted the customer, when, through which channel, and the result.", "Manual handling evidence"],
            ["6", "Open an incident or complaint case when failure creates safety, legal, quality, privacy, or material customer impact.", "Case and escalation"],
        ],
        "exceptions": ["Do not include secrets, unrestricted private documents, or internal-only notes.", "Do not claim delivery from a simulated or accepted provider response without delivery evidence.", "Prevent duplicate sends when retrying uncertain provider responses."],
        "completion": ["Required notice is delivered or a documented alternate contact is complete.", "Provider failure has an owner and resolution evidence.", "Customer-impacting exceptions are linked to a case."],
        "metric": "Notices needing attention, delivery and bounce rate, retry age, first response time, and duplicate sends", "records": "Notification and events, feedback history, provider evidence, alternate contact note, case record",
    },
]


def build_sops():
    doc = Document()
    configure(doc, "SOP LIBRARY", "Standard Operating Procedure Library")
    add_title(doc, "Standard Operating Procedure Library", "Verified application workflows for controlled business operation", "SOP LIBRARY")
    doc.add_heading("How to use this library", level=1)
    doc.add_paragraph(
        "Follow the procedure that matches the work. Stop when a prerequisite is missing or the application refuses a guarded action. Record the exception in an operational case and assign the next action. An issued procedure never authorizes a person beyond the approved authority matrix."
    )
    add_table(doc, ["Procedure", "Primary record"], [[f"{s['number']} {s['title']}", s['records'].split(',')[0]] for s in SOPS], [4.8, 1.95])
    for index, item in enumerate(SOPS):
        doc.add_page_break()
        sop_header(doc, item["number"], item["title"], item["owner"], item["purpose"], item["scope"], item["prereq"], item["records"], item["metric"])
        sop_steps(doc, item["steps"])
        sop_close(doc, item["exceptions"], item["completion"], item["metric"], item["records"])
        doc.add_heading("Approval and training", level=2)
        add_table(doc, ["Owner", "Approver", "Effective date", "Review date", "Training evidence"], [["", "", "", "", ""]], [1.25, 1.25, 1.25, 1.25, 1.75])
    path = OUT / "SOP-003-012_Standard_Operating_Procedure_Library.docx"
    doc.save(path)
    return path


def build_technical():
    doc = Document()
    configure(doc, "MAN 002", "Technical Operations and Recovery Manual")
    add_title(doc, "Technical Operations and Recovery Manual", "Safe deployment provider activation monitoring backup restore and incident response", "MAN 002")
    doc.add_heading("Manual purpose", level=1)
    doc.add_paragraph(
        "This manual governs technical changes that can affect business records, customer access, payments, shipping, documents, and recovery. It requires environment separation, evidence before activation, reversible changes, and an operational case when service or data integrity is uncertain."
    )
    doc.add_heading("Environment boundaries", level=1)
    add_table(doc, ["Environment", "Worker", "Database", "Documents", "Permitted use"], [
        ["Local", "Local runtime", "Local D1 state", "Local R2 state", "Development and automated tests"],
        ["Staging", "nexphaselabs staging", "nexphase labs staging", "nexphase documents staging", "Controlled rehearsal with synthetic data"],
        ["Production", "nexphaselabs", "nexphase labs", "nexphase documents", "Real activity only after launch approval"],
    ], [1.0, 1.35, 1.35, 1.45, 1.6])
    doc.add_paragraph(
        "The public domain currently serves the prior site. A deployment does not authorize a DNS cutover, real provider activity, production data import, or customer contact. Each action needs its own approved change record."
    )
    doc.add_heading("Access and secret control", level=1)
    bullets(doc, [
        "Use individual staff accounts and the least authority approved for the person’s responsibilities.",
        "Keep provider keys and deployment credentials in approved secret systems. Never record them in source files, Drive, forms, cases, or manuals.",
        "Protect all staging hostnames, including direct workers.dev access, and test the restriction from an unauthenticated browser.",
        "Review active users, sessions, recovery access, and leavers on the approved schedule.",
        "Revoke sessions after role, password, or employment changes when continued access is not approved.",
    ])
    doc.add_heading("Change and deployment procedure", level=1)
    sop_steps(doc, [
        ["1", "Define scope, affected workflows, data changes, owner, reviewer, rollback, and acceptance evidence.", "Approved change record"],
        ["2", "Review uncommitted work and isolate unrelated changes. Add migration and workflow tests for data changes.", "Traceable change set"],
        ["3", "Run type checking, lint, the complete test suite, production build, fresh migration replay, and foreign key check.", "Green verification evidence"],
        ["4", "Back up the target database and document bucket before a material change. Verify the database export in isolation.", "Recovery evidence"],
        ["5", "Apply migrations to staging and deploy staging. Verify health, authentication, and the affected normal and failure paths.", "Staging result"],
        ["6", "Obtain production approval. Apply migration before code when the migration is backward compatible; follow the approved sequence otherwise.", "Production change event"],
        ["7", "Run health and a non destructive smoke test. Monitor logs, provider queues, and operational cases.", "Post change verification"],
        ["8", "If acceptance fails, stop, preserve evidence, roll back code, and follow the database recovery decision. Never improvise destructive data reversal.", "Rollback or incident record"],
    ])
    doc.add_heading("Exact commands", level=1)
    doc.add_paragraph(
        "Type these verbatim. Everything in this section has been run. A second person should be able to follow it without asking a question, which is the only reason it exists: on 12 September 2026 nobody except the technology seat could deploy, roll back, or restore, and the cutover is the worst moment for that to still be true."
    )
    doc.add_paragraph(
        "Before anything else, confirm which account and environment the shell is pointed at. Never deploy from a shell whose target you have not just checked."
    )
    commands(doc, [
        "cd /path/to/nexphaselabs.net",
        "npx wrangler whoami                      # account and token scope",
        "git status --porcelain                   # must be empty before a release",
        "git log --oneline -3",
    ])

    doc.add_heading("Deploy to staging", level=2)
    doc.add_paragraph(
        "Staging deploys from main. Pushing or merging to main is itself a deployment: the Deploy staging workflow runs checks, build, staging migrations, seed, deploy, and two smoke tests. Do it while someone is watching, never as a side effect of tidying up."
    )
    commands(doc, [
        "git checkout main && git pull",
        "git merge --no-ff <branch>               # this is a deployment",
        "git push origin main                     # triggers .github/workflows/deploy-staging.yml",
        "",
        "# watch it, and read the last two steps in particular",
        "gh run watch",
        "",
        "# the same thing by hand, if GitHub is unavailable",
        "npm run check                            # typecheck, lint, tests, build",
        "npm run db:migrate:staging",
        "npm run deploy:staging                   # builds, runs the guard, deploys",
    ])
    doc.add_paragraph(
        "The deploy guard refuses to hand wrangler a build that does not target the environment named on the command line, and prints every field it checked. If it refuses, rebuild; never bypass it. The staging smoke steps require the health endpoint to report ok, the public storefront to answer 200 with a noindex header, and staff pages plus private management APIs to refuse anonymous access."
    )

    doc.add_heading("Release to production", level=2)
    doc.add_paragraph(
        "A version tag is the only path to production. There is no other way in: the workflow refuses any run whose ref is not a v tag, and the guard refuses a local deploy unless HEAD is at a v tag with a clean tree."
    )
    commands(doc, [
        "# release what is already green on staging",
        "git checkout main && git pull",
        "git tag -a v1.2.3 -m \"Launch release\"",
        "git push origin v1.2.3                   # triggers deploy-production.yml",
        "gh run watch",
        "",
        "# break glass only: GitHub unavailable, HEAD at the tag, clean tree",
        "NX_CONFIRM_PRODUCTION=yes npm run deploy",
    ])

    doc.add_heading("Roll back", level=2)
    doc.add_paragraph(
        "Code rolls back. Migrations do not: they are written to be backward compatible so the previous worker keeps running against the new schema. If a migration has made the data itself wrong, this is an incident and the database decision below applies, not a rollback."
    )
    commands(doc, [
        "npx wrangler deployments list                       # production",
        "npx wrangler deployments list --env staging",
        "npx wrangler rollback --message \"Reason, case number\"",
        "npx wrangler rollback <deployment-id> --env staging",
        "",
        "# prove the rollback landed",
        "curl -fsS https://nexphaselabs.net/api/health",
        "npx wrangler tail                                   # live log, Ctrl-C to stop",
    ])
    doc.add_paragraph(
        "Record the rollback in an operational case with the deployment identifier, the reason, who approved it, and what the health check said afterwards."
    )

    doc.add_heading("Back up and restore", level=2)
    doc.add_paragraph(
        "One command runs the whole rehearsal, times every phase, and writes a report with a witness block. It checks credentials before it touches anything. Without R2 credentials the document half is skipped and the result is reported incomplete, because a database restored without its evidence files is not a recovery."
    )
    commands(doc, [
        "export CLOUDFLARE_API_TOKEN=...          # D1 Edit + R2 Edit, correct account",
        "export R2_ACCESS_KEY_ID=...              # R2 S3 credentials",
        "export R2_SECRET_ACCESS_KEY=...",
        "",
        "npm run ops:rehearsal:staging -- /absolute/private/recovery-dir",
        "",
        "# production, inside an approved maintenance window only",
        "NEXPHASE_CONFIRM_PRODUCTION_BACKUP=yes \\",
        "  npm run ops:rehearsal:production /absolute/private/recovery-dir",
        "",
        "# the individual steps, if you need them separately",
        "npm run ops:backup:staging -- /absolute/private/recovery-dir",
        "npm run ops:restore:verify -- /absolute/private/recovery-dir/<export>.sql",
    ])
    doc.add_paragraph(
        "Authentication error 10000 means the token was rejected, not that the database is missing. Check that the token exists and has not expired, that it carries D1 Edit and R2 Edit, that it is scoped to the account holding the database, and that the shell really has it exported. The rehearsal reports this for you."
    )
    doc.add_paragraph(
        "Cloudflare Time Travel restores a D1 database to a point in the last thirty days. It is a data decision, not a technical one: it discards everything written since that timestamp. It requires the administrator's approval and an incident case, and the export above is taken first."
    )
    commands(doc, [
        "npx wrangler d1 time-travel info nexphase-labs",
        "npx wrangler d1 time-travel restore nexphase-labs --timestamp <ISO-8601>",
    ])

    doc.add_heading("Staging access", level=2)
    doc.add_paragraph(
        "Public staging pages are intentionally open for review and carry noindex on every response. Staff pages and private management APIs remain behind the application's staff authentication. The deploy fails if the public/private boundary or noindex header changes."
    )
    commands(doc, [
        "# prove it from outside, unauthenticated",
        "curl -s -o /dev/null -w '%{http_code}\\n' https://<staging-host>/     # expect 200",
        "curl -sI https://<staging-host>/ | grep -i x-robots-tag              # expect noindex",
        "curl -sI https://<staging-host>/manage | grep -i '^location: /staff/sign-in'",
        "curl -s -o /dev/null -w '%{http_code}\\n' -X POST -H 'Origin: https://<staging-host>' https://<staging-host>/api/manage/notifications  # expect 401",
    ])

    doc.add_heading("Where the credentials live", level=1)
    doc.add_paragraph(
        "Locations and who can obtain them. No value is ever written here, in any other manual, in a case, or in source. If a credential must be replaced, rotate it at the provider and set it again through the command that owns it."
    )
    add_table(doc, ["Credential", "Where it lives", "Who can obtain it", "Used by"], [
        ["Cloudflare account access", "Cloudflare dashboard login with the account owner's own 2FA", "Business and Systems lead", "Every deploy, D1, R2, logs"],
        ["CLOUDFLARE_API_TOKEN", "GitHub repository secret, and the operator's own shell for manual runs", "Business and Systems lead, from the Cloudflare dashboard", "Both deploy workflows, D1 export"],
        ["CLOUDFLARE_ACCOUNT_ID", "GitHub repository secret and wrangler.jsonc (not secret)", "Anyone with repository access", "Deploy workflows"],
        ["R2 S3 access key and secret", "Cloudflare R2 API tokens page, held by the operator for a rehearsal", "Business and Systems lead", "Recovery rehearsal only"],
        ["Worker secrets", "Cloudflare Workers secret store, per environment, set by wrangler secret put", "Business and Systems lead", "Payments, email, Shippo, digest"],
        ["Google Workspace sending credentials", "Workspace admin console and the Worker secret store", "Business and Systems lead", "Customer email"],
        ["Shippo API key", "Shippo dashboard and the Worker secret store", "Business and Systems lead", "Rates, labels, tracking"],
        ["Payment provider credentials", "Provider dashboard and the Worker secret store", "Business and Systems lead with banking owner approval", "Checkout and refunds"],
        ["Domain and DNS", "Registrar and Cloudflare DNS", "Business and Systems lead", "Cutover"],
    ], [1.5, 2.2, 1.6, 1.45])
    doc.add_paragraph(
        "A second person cannot obtain these today. Until a named fallback exists with their own access, a single unavailable person is a single point of failure for deployment and recovery. Record the fallback and their access path here when it is agreed."
    )

    doc.add_heading("Provider activation", level=1)
    add_table(doc, ["Provider area", "Required staging evidence", "Production gate"], [
        ["USPS UPS FedEx through Shippo", "Named origins, approved parcels and services, sandbox rates, one label, refund reconciliation, replacement after void, tracking, dangerous goods decision", "Carrier accounts, each origin, service policy, return and handoff procedure approved"],
        ["Payment", "Invoice or payment creation, webhook or reconciliation, failure recovery, refund, settlement match", "Banking owner approves account, credentials, refund and reconciliation policy"],
        ["Tax", "Representative address calculations, exemptions where applicable, error behavior, report mapping", "CPA or tax adviser approves nexus, taxability, filing owner, and configuration"],
        ["Email", "Approved sender, inbox delivery, bounce or failure, durable retry, manual recovery", "Domain authentication, recipients, retention, and support ownership approved"],
    ], [1.45, 3.05, 2.25])
    doc.add_paragraph(
        "USPS is implemented through the existing Shippo rate and label boundary. Staff select a named ship from location before quoting, and the origin remains in quote and label history. Enable only approved USPS services and accurate packed dimensions. Record delivery only from carrier evidence. USPS commercial Ground Advantage rules require dimensions in applicable cases, and hazardous material rules remain the shipper’s responsibility."
    )
    doc.add_paragraph(
        "Label cancellation uses one durable refund claim. The application sends one Shippo refund request, blocks shipment replacement and order cancellation while the outcome is pending or uncertain, and reconciles through a read only refund lookup. If the refund response is interrupted before an identifier returns, Check refund status searches existing refunds by transaction without creating another. A replacement label is allowed only after the prior label is confirmed voided, and every label remains in order history."
    )
    doc.add_heading("Health monitoring", level=1)
    bullets(doc, [
        "Monitor the health endpoint for complete database schema access and document store availability without exposing business data.",
        "Monitor failed customer notices, payment attempts, shipping quote and label errors, and critical or overdue operational cases.",
        "A healthy endpoint proves technical access at that moment. It does not prove inventory, provider settlement, legal approval, backup recovery, or launch readiness.",
        "Open an incident when health is unavailable, schema probes fail, documents cannot be retrieved, or state may be incomplete.",
    ])
    doc.add_heading("Database backup and restore rehearsal", level=1)
    sop_steps(doc, [
        ["1", "Create a dedicated encrypted backup directory outside the repository.", "Restricted destination"],
        ["2", "Run the staging D1 backup script with the absolute directory path.", "SQL export and checksum manifest"],
        ["3", "Run the restore verification script against the export.", "Integrity ok, no foreign key violations, required tables present"],
        ["4", "Copy the R2 bucket through approved read only S3 credentials to the dated recovery directory and save an object manifest.", "Complete object copy"],
        ["5", "Restore to a new isolated rehearsal bucket and compare count, bytes, and representative object checksums.", "R2 recovery evidence"],
        ["6", "Verify authorized and unauthorized document retrieval against disposable isolated bindings.", "Access behavior evidence"],
        ["7", "Record elapsed time, results, reviewer, exceptions, and the next rehearsal date in Operating Controls.", "Reviewed continuity control"],
    ])
    doc.add_paragraph(
        "Use node scripts d1 backup with staging and the private output directory. Production export requires the explicit production confirmation environment value and an approved maintenance window. A rehearsal must never import into the source database or overwrite the source bucket."
    )
    doc.add_heading("Incident response", level=1)
    sop_steps(doc, [
        ["1", "Confirm the signal and open an incident case with severity, owner, due date, and known systems or records.", "Incident case"],
        ["2", "Contain high or critical impact. Stop affected provider, deployment, shipment, release, or customer action when needed.", "Containment record"],
        ["3", "Preserve logs, timestamps, configuration identifiers, checksums, and affected record numbers without copying secrets into the case.", "Evidence link"],
        ["4", "Notify authorized operations, quality, commercial, legal, or privacy owners according to impact.", "Communication history"],
        ["5", "Recover through the approved rollback, provider recovery, or isolated restore path. Verify data and access before reopening service.", "Recovery evidence"],
        ["6", "Record root cause, corrective and preventive actions, then perform an effectiveness check.", "CAPA record"],
        ["7", "Administrator reviews closure. Update controls, monitoring, tests, and manuals when the corrective action changes operation.", "Closed incident and controlled revision"],
    ])
    doc.add_heading("Release rehearsal", level=1)
    doc.add_paragraph(
        "Run the named team through supplier qualification, PO, receipt, quarantine, testing, release, inventory movement, catalog, checkout, payment, order ownership, pick and pack, USPS UPS and FedEx comparison, label, shipment, delivery evidence, notification, return, refund, period reports, close, backup, and one operational case. Inject one failure at each provider boundary and record recovery."
    )
    add_table(doc, ["Acceptance", "Evidence"], [
        ["All automated checks pass", "Type check, lint, full tests, build, migration replay"],
        ["No unresolved priority zero defect", "Defect and case review"],
        ["Every launch critical control reviewed", "Operating Controls export or screenshots and evidence links"],
        ["Recovery completes within approved targets", "D1 and R2 restore evidence with elapsed time"],
        ["Named operators complete procedures", "Training and walkthrough records"],
        ["Authorized members approve launch", "Signed decision with scope and residual risks"],
    ], [3.3, 3.45])
    doc.add_heading("References", level=1)
    bullets(doc, [
        "Cloudflare D1 import and export data  https://developers.cloudflare.com/d1/best-practices/import-export-data/",
        "Cloudflare D1 Time Travel  https://developers.cloudflare.com/d1/reference/time-travel/",
        "Cloudflare R2 download objects  https://developers.cloudflare.com/r2/objects/download-objects/",
        "Shippo carrier accounts and rates  https://docs.goshippo.com/api-reference/carrier-accounts/list-all-carrier-accounts",
        "Shippo service levels  https://docs.goshippo.com/api-reference/service-levels",
        "Shippo create a refund  https://docs.goshippo.com/api-reference/refunds/create-a-refund",
        "Shippo list all refunds  https://docs.goshippo.com/api-reference/refunds/list-all-refunds",
        "Shippo retrieve a refund  https://docs.goshippo.com/api-reference/refunds/retrieve-a-refund",
        "USPS Ground Advantage  https://www.usps.com/ship/ground-advantage.htm",
    ])
    doc.add_heading("Approval record", level=1)
    add_table(doc, ["Role", "Name", "Decision", "Date", "Evidence"], [["Technical owner", "", "", "", ""], ["Operations reviewer", "", "", "", ""], ["Quality reviewer", "", "", "", ""], ["Production approver", "", "", "", ""]], [1.45, 1.25, 1.15, 1.0, 1.9])
    path = OUT / "MAN-002_Technical_Operations_and_Recovery_Manual.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    paths = [build_master(), build_sops(), build_technical()]
    for path in paths:
        print(path)
